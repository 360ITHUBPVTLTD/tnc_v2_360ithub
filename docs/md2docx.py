"""Markdown -> .docx for the CI/CD guide. Handles the subset used in that file:
headings, paragraphs, bold/italic/inline-code, bullet + numbered lists, tables,
fenced code blocks, blockquotes (including lists inside them) and page breaks."""
import re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CODE_FONT = "Consolas"
CODE_BG = "F2F2F2"
QUOTE_BG = "FFF6E5"


def shade(el, color):
    pr = el.get_or_add_pPr() if hasattr(el, "get_or_add_pPr") else el
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    pr.append(shd)


def left_bar(p, color="E0A030"):
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    b = OxmlElement("w:left")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "18")
    b.set(qn("w:space"), "8"); b.set(qn("w:color"), color)
    borders.append(b)
    pPr.append(borders)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<![\w*])\*[^*\n]+\*(?![\w*]))")


def add_inline(p, text, base_size=None, bold_all=False, italic_all=False):
    """Recursive so nested marks like **`code`** render as bold code, not literal
    backticks."""
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            add_inline(p, tok[2:-2], base_size, True, italic_all)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = p.add_run(tok[1:-1])
            r.font.name = CODE_FONT
            r.font.size = base_size or Pt(9.5)
            r.font.color.rgb = RGBColor(0xA3, 0x11, 0x5E)
            if bold_all:
                r.bold = True
            if italic_all:
                r.italic = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            add_inline(p, tok[1:-1], base_size, bold_all, True)
        else:
            r = p.add_run(tok)
            if bold_all:
                r.bold = True
            if italic_all:
                r.italic = True
            if base_size:
                r.font.size = base_size
    return p


def add_code(doc, lines, indent=0.0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.18 + indent); pf.right_indent = Inches(0.1)
    pf.space_before = Pt(6); pf.space_after = Pt(8)
    pf.line_spacing = 1.0
    # An ASCII diagram split across a page break is unreadable. Only for blocks
    # that can actually fit on a page - a 200-line one would leave a huge gap.
    if len(lines) <= 45:
        pf.keep_together = True
    shade(p._p, CODE_BG)
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        r = p.add_run(line.rstrip("\n"))
        r.font.name = CODE_FONT; r.font.size = Pt(8.5)
    return p


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.autofit = True
    for i, cell_text in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        add_inline(c.paragraphs[0], cell_text, base_size=Pt(9.5), bold_all=True)
        shade(c._tc.get_or_add_tcPr(), "DCE6F1")
    for row in body:
        cells = t.add_row().cells
        for i, cell_text in enumerate(row[:len(header)]):
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], cell_text, base_size=Pt(9.5))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_sep(line):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def render(doc, lines, indent=0.0, quoted=False, first_h1_seen=None):
    if first_h1_seen is None:
        first_h1_seen = [False]
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()

        if not s:
            i += 1; continue

        # fenced code
        if s.startswith("```"):
            j = i + 1; buf = []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                buf.append(lines[j]); j += 1
            add_code(doc, buf, indent)
            i = j + 1; continue

        # blockquote
        if s.startswith(">"):
            j = i; buf = []
            while j < len(lines) and lines[j].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[j])); j += 1
            start = len(doc.paragraphs)
            render(doc, buf, indent + 0.25, True, first_h1_seen)
            for p in doc.paragraphs[start:]:
                shade(p._p, QUOTE_BG)
                left_bar(p)
                p.paragraph_format.left_indent = Inches(0.3 + indent)
            i = j; continue

        # table
        if s.startswith("|") and i + 1 < len(lines) and is_sep(lines[i + 1]):
            rows = [split_row(s)]
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append(split_row(lines[j])); j += 1
            add_table(doc, rows)
            i = j; continue

        # indented code block (4 spaces), the other markdown code form
        if not quoted and line.startswith("    ") and line[4:5] not in ("", " "):
            j = i
            buf = []
            while j < len(lines):
                cur = lines[j]
                if cur.startswith("    "):
                    buf.append(cur[4:])
                    j += 1
                elif not cur.strip():
                    # a blank line continues the block only if more indent follows
                    k = j
                    while k < len(lines) and not lines[k].strip():
                        k += 1
                    if k < len(lines) and lines[k].startswith("    "):
                        buf.extend([""] * (k - j))
                        j = k
                    else:
                        break
                else:
                    break
            add_code(doc, buf, indent)
            i = j
            continue

        # horizontal rule
        if re.fullmatch(r"-{3,}", s):
            i += 1; continue

        # headings
        m = re.match(r"^(#{1,5})\s+(.*)$", s)
        if m:
            level, text = len(m.group(1)), m.group(2)
            if level == 1 and not first_h1_seen[0]:
                first_h1_seen[0] = True
                p = doc.add_paragraph(); p.style = doc.styles["Title"]
                add_inline(p, text)
            elif level == 1:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                p = doc.add_heading("", level=1); add_inline(p, text)
            else:
                if quoted:
                    p = doc.add_paragraph()
                    add_inline(p, text, base_size=Pt(11), bold_all=True)
                else:
                    p = doc.add_heading("", level=min(level, 4))
                    add_inline(p, text)
            i += 1; continue

        # bullet list
        if re.match(r"^[-*]\s+", s):
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                txt = re.sub(r"^\s*[-*]\s+", "", lines[i].strip())
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.left_indent = Inches(0.35 + indent)
                p.paragraph_format.space_after = Pt(2)
                add_inline(p, txt)
                i += 1
            continue

        # numbered list
        if re.match(r"^\d+\.\s+", s):
            while i < len(lines) and re.match(r"^\s*\d+\.\s+", lines[i]):
                txt = re.sub(r"^\s*\d+\.\s+", "", lines[i].strip())
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.left_indent = Inches(0.35 + indent)
                p.paragraph_format.space_after = Pt(2)
                add_inline(p, txt)
                i += 1
            continue

        # paragraph (join continuation lines)
        buf = [s]; j = i + 1
        while j < len(lines):
            nxt = lines[j].strip()
            if (not nxt or nxt.startswith(("#", ">", "|", "```", "- ", "* "))
                    or re.match(r"^\d+\.\s", nxt) or re.fullmatch(r"-{3,}", nxt)):
                break
            buf.append(nxt); j += 1
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(indent)
        p.paragraph_format.space_after = Pt(6)
        add_inline(p, " ".join(buf))
        i = j


def add_page_numbers(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = 1
    r = p.add_run("Page ")
    r.font.size = Pt(9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)


def main(src, dst):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.8)
        sec.top_margin = sec.bottom_margin = Inches(0.7)
    lines = open(src, encoding="utf-8").read().splitlines()
    render(doc, lines)
    add_page_numbers(doc)
    doc.save(dst)
    print(f"wrote {dst}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
