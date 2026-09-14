import frappe
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from frappe.utils import getdate, add_days, get_first_day, get_last_day, today, formatdate

def debug_info():
	if not frappe.db.exists("Report", "Monthly Teacher Task Summary Report"):
		doc = frappe.get_doc({
			"doctype": "Report",
			"report_name": "Monthly Teacher Task Summary Report",
			"ref_doctype": "Task",
			"report_type": "Script Report",
			"is_standard": "Yes",
			"module": "TNC v2",
			"roles": [
				{"role": "System Manager"},
				{"role": "TNC Teachers"}
			]
		})
		doc.insert(ignore_permissions=True)
		frappe.db.commit()
		print("REPORT CREATED SUCCESSFUL")
	else:
		print("REPORT ALREADY EXISTS")

def check_email_logs():
	emails = frappe.db.get_all("Email Queue", fields=["name", "message", "status", "sender", "creation"], order_by="creation desc", limit=5)
	for e in emails:
		print("EMAIL:", e.name, "| sender:", e.sender, "| status:", e.status, "| creation:", e.creation)

@frappe.whitelist()
def download_excel_report(**kwargs):
	"""
	Generates and returns the styled Excel report for download.
	"""
	filters = {}
	for k, v in kwargs.items():
		try:
			filters[k] = frappe.parse_json(v)
		except Exception:
			filters[k] = v

	from tnc_v2_360ithub.tnc_v2.report.monthly_teacher_task_summary_report.monthly_teacher_task_summary_report import get_data
	tasks = get_data(filters)

	date_range = filters.get("date_range")
	if date_range and len(date_range) == 2:
		from_date = getdate(date_range[0])
		to_date = getdate(date_range[1])
		period_label = f"{from_date.strftime('%d-%m-%Y')} to {to_date.strftime('%d-%m-%Y')}"
	else:
		period_label = "All Time"

	user_name = "All Users"
	if filters.get("task_owner"):
		user_name = frappe.db.get_value("User", filters.get("task_owner"), "full_name") or filters.get("task_owner")

	excel_content = generate_excel_report(tasks, user_name, period_label, filters)

	frappe.local.response.filename = f"Monthly_Task_Summary_Report_{user_name.replace(' ', '_')}.xlsx"
	frappe.local.response.filecontent = excel_content
	frappe.local.response.type = "binary"

def send_monthly_teacher_task_summary_reports():
	"""
	Automatically run on 1st of every month at 7:00 AM.
	Generates and emails a Monthly Task Summary Report to each employee/teacher individually.
	"""
	today_date = getdate(today())
	first_day_of_this_month = get_first_day(today_date)
	last_day_of_prev_month = add_days(first_day_of_this_month, -1)
	first_day_of_prev_month = get_first_day(last_day_of_prev_month)
	
	from_date = first_day_of_prev_month.strftime("%Y-%m-%d")
	to_date = last_day_of_prev_month.strftime("%Y-%m-%d")

	send_reports(from_date, to_date)

def generate_excel_report(tasks, user_name, period_name, filters=None):
	"""
	Generates a highly styled, professional Excel worksheet for tasks.
	"""
	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = "Task Summary"
	
	ws.views.sheetView[0].showGridLines = True
	
	font_family = "Segoe UI"
	title_font = Font(name=font_family, size=15, bold=True, color="1B365D")
	header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
	data_font = Font(name=font_family, size=11, color="2C3E50")
	bold_data_font = Font(name=font_family, size=11, bold=True, color="2C3E50")
	filter_label_font = Font(name=font_family, size=10, bold=True, color="475569")
	filter_val_font = Font(name=font_family, size=10, color="1E293B")
	
	header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
	zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
	
	thin_border = Border(
		left=Side(style='thin', color='E2E8F0'),
		right=Side(style='thin', color='E2E8F0'),
		top=Side(style='thin', color='E2E8F0'),
		bottom=Side(style='thin', color='E2E8F0')
	)
	
	# Title Row
	ws.merge_cells("A1:I1")
	ws["A1"] = f"Monthly Task Summary Report - {user_name}"
	ws["A1"].font = title_font
	ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
	ws.row_dimensions[1].height = 35
	
	# Period Row
	ws.merge_cells("A2:I2")
	ws["A2"] = f"Period: {period_name}"
	ws["A2"].font = Font(name=font_family, size=11, italic=True, color="64748B")
	ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
	ws.row_dimensions[2].height = 20
	
	# Active Filters Block
	curr_row = 4
	if filters:
		active_filters = []
		if filters.get("task_owner"):
			active_filters.append(("Task Owner", filters.get("task_owner")))
		if filters.get("department"):
			active_filters.append(("Department", filters.get("department")))
		if filters.get("status"):
			active_filters.append(("Status", filters.get("status")))
		if filters.get("priority"):
			active_filters.append(("Priority", filters.get("priority")))
		if filters.get("task_reporter"):
			active_filters.append(("Task Reporter", filters.get("task_reporter")))
		if filters.get("subject"):
			active_filters.append(("Subject", filters.get("subject")))

		if active_filters:
			# Header for filters section
			ws.cell(row=curr_row, column=1, value="Selected Filters:").font = Font(name=font_family, size=10, bold=True, color="1B365D", underline="single")
			curr_row += 1
			
			for i in range(0, len(active_filters), 2):
				ws.row_dimensions[curr_row].height = 18
				lbl1, val1 = active_filters[i]
				ws.cell(row=curr_row, column=1, value=lbl1).font = filter_label_font
				ws.cell(row=curr_row, column=2, value=val1).font = filter_val_font
				
				if i + 1 < len(active_filters):
					lbl2, val2 = active_filters[i+1]
					ws.cell(row=curr_row, column=4, value=lbl2).font = filter_label_font
					ws.cell(row=curr_row, column=5, value=val2).font = filter_val_font
				curr_row += 1
			curr_row += 1 # Empty spacing row
			
	# Table Headers
	header_row = curr_row
	headers = ["Task ID", "Subject", "Task Reporter", "Task Owner Name", "Status", "Priority", "Expected Start Date", "Expected End Date", "Overdue Days"]
	ws.row_dimensions[header_row].height = 26
	for col_idx, text in enumerate(headers, 1):
		cell = ws.cell(row=header_row, column=col_idx, value=text)
		cell.font = header_font
		cell.fill = header_fill
		cell.alignment = Alignment(horizontal="center" if col_idx != 2 else "left", vertical="center")
		cell.border = thin_border
		
	# Freeze rows up to header row
	ws.freeze_panes = f"A{header_row + 1}"
	
	row_num = header_row + 1
	if tasks:
		for idx, t in enumerate(tasks):
			exp_start_str = formatdate(t.get("exp_start_date"), "dd-mm-yyyy") if t.get("exp_start_date") else "-"
			exp_end_str = formatdate(t.get("exp_end_date"), "dd-mm-yyyy") if t.get("exp_end_date") else "-"
			overdue_days = t.get("overdue_days", 0)

			task_id_val = t.get("task_id") or t.get("name")
			row_data = [
				task_id_val,
				t["subject"] or "-",
				t["task_reporter_name"] or t["task_reporter"] or "-",
				t.get("task_owner_name") or t.get("task_owner") or "-",
				t["status"],
				t["priority"],
				exp_start_str,
				exp_end_str,
				overdue_days
			]

			subject_lines = max(1, -(-len(str(row_data[1])) // 56))
			ws.row_dimensions[row_num].height = 22 if subject_lines <= 1 else 16 * subject_lines
			for col_idx, val in enumerate(row_data, 1):
				cell = ws.cell(row=row_num, column=col_idx, value=val)
				cell.font = data_font
				cell.border = thin_border

				if col_idx in [1, 5, 6, 7, 8]:
					cell.alignment = Alignment(horizontal="center", vertical="center")
				elif col_idx == 9:
					cell.alignment = Alignment(horizontal="center", vertical="center")
					cell.font = bold_data_font
				elif col_idx == 2:
					cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
				else:
					cell.alignment = Alignment(horizontal="left", vertical="center")

				if col_idx == 1 and task_id_val:
					cell.hyperlink = f"{frappe.utils.get_url()}/app/task/{task_id_val}"
					cell.font = Font(name=font_family, size=11, color="2563EB", underline="single")

				if idx % 2 == 1:
					cell.fill = zebra_fill
			row_num += 1
	else:
		ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=9)
		cell = ws.cell(row=row_num, column=1, value="No Tasks Found")
		cell.font = Font(name=font_family, size=11, italic=True, color="94A3B8")
		cell.alignment = Alignment(horizontal="center", vertical="center")
		cell.border = thin_border
		ws.row_dimensions[row_num].height = 35
		
	for col in ws.columns:
		max_len = 0
		for cell in col:
			val_str = str(cell.value or '')
			if cell.coordinate in ws.merged_cells:
				continue
			if len(val_str) > max_len:
				max_len = len(val_str)
		col_letter = openpyxl.utils.get_column_letter(col[0].column)
		ws.column_dimensions[col_letter].width = min(max(max_len + 4, 13), 60)
		
	out = io.BytesIO()
	wb.save(out)
	out.seek(0)
	return out.getvalue()

def send_reports(from_date, to_date):
	try:
		from_date_val = getdate(from_date)
		to_date_val = getdate(to_date)
		month_name = from_date_val.strftime("%B %Y")
		from_date_str = from_date_val.strftime("%d-%m-%y")
		to_date_str = to_date_val.strftime("%d-%m-%y")
		period_label = f"{from_date_str} to {to_date_str}"
		
		users_to_report = {}
		
		employees = frappe.db.get_all(
			"Employee",
			filters={"status": "Active", "user_id": ["is", "set"]},
			fields=["employee_name", "user_id"]
		)
		for emp in employees:
			if emp.user_id:
				users_to_report[emp.user_id.lower().strip()] = {
					"name": emp.employee_name,
					"email": emp.user_id.strip(),
					"type": "Employee"
				}
				
		teachers = frappe.db.get_all(
			"Teacher",
			filters={"status": "Active", "email": ["is", "set"]},
			fields=["full_name", "email"]
		)
		for teacher in teachers:
			if teacher.email:
				email_lower = teacher.email.lower().strip()
				if email_lower not in users_to_report:
					users_to_report[email_lower] = {
						"name": teacher.full_name or teacher.name,
						"email": teacher.email.strip(),
						"type": "Teacher"
					}

		for email, user_info in users_to_report.items():
			tasks = frappe.db.get_all(
				"Task",
				filters={
					"task_owner": user_info["email"],
					"exp_start_date": ["between", [from_date, to_date]]
				},
				fields=["name", "subject", "task_owner_name", "task_reporter_name", "task_reporter", "status", "priority", "exp_start_date", "exp_end_date"]
			)
			
			html_message = f"""
			<div style="font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; padding: 30px; max-width: 600px; margin: auto; background-color: #FFFFFF; border: 1px solid #E2E8F0; border-radius: 12px; box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.05), 0 2px 4px -1px rgba(0, 0, 0, 0.03);">
				<div style="border-bottom: 2px solid #F1F5F9; padding-bottom: 20px; margin-bottom: 25px; display: flex; align-items: center; justify-content: space-between;">
					<div>
						<h2 style="color: #1E3A8A; margin: 0; font-size: 20px; font-weight: 800; letter-spacing: -0.025em;">Monthly Task Summary</h2>
						<p style="color: #64748B; margin: 5px 0 0 0; font-size: 14px; font-weight: 500;">Period: <span style="color: #0F172A; font-weight: 600;">{period_label}</span></p>
					</div>
					<div style="text-align: right;">
						<span style="background-color: #EFF6FF; color: #1D4ED8; padding: 6px 12px; border-radius: 8px; font-size: 12px; font-weight: 700; border: 1px solid #DBEAFE;">{month_name}</span>
					</div>
				</div>
				
				<div style="margin-bottom: 25px; font-size: 15px; line-height: 1.6; color: #334155;">
					<p>Dear <strong>{user_info['name']}</strong>,</p>
					<p>Please find attached your Monthly Task Summary Report for the period of <strong>{period_label}</strong>.</p>
					<p>The attached Excel spreadsheet contains detailed information regarding your assigned tasks, including their subjects, reporters, statuses, priorities, expected start/end dates, and overdue days.</p>
					<p>If you have any questions or require modifications to your tasks, please contact your department head or administrator.</p>
				</div>
				
				<div style="margin-top: 30px; border-top: 1px solid #F1F5F9; padding-top: 20px; font-size: 13px; color: #475569; line-height: 1.5;">
					<p style="margin: 0;">Best regards,</p>
					<p style="margin: 5px 0 0 0; font-weight: 700; color: #1E3A8A; font-size: 14px;">TNC Institute Management System</p>
					<p style="margin: 2px 0 0 0; color: #94A3B8; font-size: 12px;">Automated Report System</p>
				</div>
			</div>
			"""
			
			subject = f"Monthly Task Summary Report - {user_info['name']} - {month_name}"
			excel_content = generate_excel_report(tasks, user_info["name"], period_label)
			
			attachments = [{
				"fname": f"Monthly_Task_Summary_{user_info['name'].replace(' ', '_')}_{month_name.replace(' ', '_')}.xlsx",
				"fcontent": excel_content
			}]
			
			try:
				frappe.sendmail(
					recipients=[user_info["email"]],
					subject=subject,
					message=html_message,
					attachments=attachments,
					delayed=False
				)
			except frappe.OutgoingEmailError:
				frappe.logger().warning(f"Could not send Monthly Task Summary to {user_info['email']}: No default outgoing Email Account setup.")
		
		frappe.db.commit()
		print("EMAILS PROCESSED SUCCESSFUL")
			
	except Exception as e:
		frappe.log_error(frappe.get_traceback(), "Monthly Teacher Task Summary Report Schedular Error")
