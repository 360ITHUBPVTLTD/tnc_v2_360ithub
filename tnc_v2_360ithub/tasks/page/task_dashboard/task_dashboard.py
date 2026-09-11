import frappe
from frappe import _
from frappe.utils import getdate, nowdate, cint, get_url

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


@frappe.whitelist()
def get_filter_options():
    return {
        "statuses": ["Open", "Working", "Pending Review", "Completed", "Cancelled", "Overdue"],
        "priorities": ["Low", "Medium", "High", "Urgent"],
        "owners": frappe.get_all(
            "User",
            filters={"enabled": 1},
            fields=["name", "full_name"],
            order_by="full_name asc"
        ),
        "companies": frappe.get_all(
            "Company",
            pluck="name",
            order_by="name asc"
        )
    }


@frappe.whitelist()
def get_dashboard_data(filters=None, start=0, page_length=20, sort_field="creation", sort_order="desc"):
    filters = frappe.parse_json(filters) if filters else {}
    start = cint(start)
    page_length = cint(page_length)

    base_filters = filters.copy()
    if "card_filter" in base_filters:
        del base_filters["card_filter"]
    if "status" in base_filters:
        del base_filters["status"]

    base_conditions = get_conditions(base_filters)
    conditions = get_conditions(filters)

    kpis = get_kpis(base_conditions)
    tasks = get_tasks(conditions, start, page_length, sort_field, sort_order, filters)
    total_count = get_total_count(conditions)

    return {
        "kpis": kpis,
        "tasks": tasks,
        "total_count": total_count
    }


def get_conditions(filters):
    conditions = []

    if filters.get("search"):
        search = frappe.db.escape(f"%{filters.get('search')}%")
        conditions.append(f"(t.name LIKE {search} OR t.subject LIKE {search})")

    if filters.get("status"):
        statuses = filters.get("status")
        if isinstance(statuses, str):
            statuses = [statuses]
            
        status_conditions = []
        regular_statuses = []
        for s in statuses:
            if s == "Overdue":
                status_conditions.append("t.status = 'Overdue'")
            else:
                regular_statuses.append(frappe.db.escape(s))
                
        if regular_statuses:
            status_conditions.append(f"t.status IN ({', '.join(regular_statuses)})")
            
        if status_conditions:
            conditions.append(f"({' OR '.join(status_conditions)})")

    if filters.get("priority"):
        priorities = filters.get("priority")
        if isinstance(priorities, str): priorities = [priorities]
        if priorities:
            conditions.append(f"t.priority IN ({', '.join(frappe.db.escape(p) for p in priorities)})")

    if filters.get("task_owner"):
        owners = filters.get("task_owner")
        if isinstance(owners, str): owners = [owners]
        if owners:
            conditions.append(f"t.task_owner IN ({', '.join(frappe.db.escape(o) for o in owners)})")

    if filters.get("company"):
        companies = filters.get("company")
        if isinstance(companies, str): companies = [companies]
        if companies:
            conditions.append(f"t.company IN ({', '.join(frappe.db.escape(c) for c in companies)})")

    if filters.get("from_date"):
        conditions.append(f"t.exp_end_date >= {frappe.db.escape(filters.get('from_date'))}")

    if filters.get("to_date"):
        conditions.append(f"t.exp_end_date <= {frappe.db.escape(filters.get('to_date'))}")

    if filters.get("card_filter"):
        f = filters.get("card_filter")
        if f == "Overdue":
            conditions.append("t.status = 'Overdue'")
        elif f == "Today":
            conditions.append("(t.status NOT IN ('Completed', 'Cancelled', 'Overdue') AND t.exp_end_date = CURDATE())")
        elif f == "Future":
            conditions.append("(t.status NOT IN ('Completed', 'Cancelled') AND t.exp_end_date > CURDATE())")
        elif f == "Pending Review":
            conditions.append("t.status = 'Pending Review'")

    return " AND ".join(conditions) if conditions else "1=1"


def get_kpis(conditions):
    data = frappe.db.sql(f"""
        SELECT
            SUM(CASE WHEN t.status = 'Overdue' THEN 1 ELSE 0 END) AS overdue,
            SUM(CASE WHEN t.status NOT IN ('Completed', 'Cancelled', 'Overdue') AND t.exp_end_date = CURDATE() THEN 1 ELSE 0 END) AS today,
            SUM(CASE WHEN t.status NOT IN ('Completed', 'Cancelled') AND t.exp_end_date > CURDATE() THEN 1 ELSE 0 END) AS future,
            SUM(CASE WHEN t.status='Pending Review' THEN 1 ELSE 0 END) AS pending_review
        FROM `tabTask` t
        WHERE {conditions}
    """, as_dict=True)[0]

    return {
        "overdue": data.overdue or 0,
        "today": data.today or 0,
        "future": data.future or 0,
        "pending_review": data.pending_review or 0
    }


def get_tasks(conditions, start, page_length, sort_field="creation", sort_order="desc", filters=None):
    valid_sort_fields = {
        "task_owner_name": "t.task_owner_name",
        "task_reporter_name": "t.task_reporter_name",
        "exp_start_date": "t.exp_start_date",
        "exp_end_date": "t.exp_end_date",
        "status": "t.status",
        "priority": "t.priority",
        "creation": "t.creation"
    }

    db_sort_field = valid_sort_fields.get(sort_field, "t.creation")
    db_sort_order = "ASC" if sort_order.lower() == "asc" else "DESC"

    if sort_field == "creation":
        order_by_clause = f"""
            CASE
                WHEN t.status = 'Overdue' THEN 0
                ELSE 1
            END,
            t.creation DESC
        """
    elif sort_field == "priority":
        order_by_clause = f"""
            CASE
                WHEN t.priority = 'Urgent' THEN 4
                WHEN t.priority = 'High' THEN 3
                WHEN t.priority = 'Medium' THEN 2
                WHEN t.priority = 'Low' THEN 1
                ELSE 0
            END {db_sort_order}
        """
    elif sort_field == "status":
        # Custom logic for Status Order
        order_by_clause = f"""
            CASE
                WHEN t.status = 'Open' THEN 1
                WHEN t.status = 'Working' THEN 2
                WHEN t.status = 'Pending Review' THEN 3
                WHEN t.status = 'Overdue' THEN 4
                WHEN t.status = 'Completed' THEN 5
                WHEN t.status = 'Cancelled' THEN 6
                ELSE 7
            END {db_sort_order}
        """
    else:
        order_by_clause = f"{db_sort_field} {db_sort_order}"

    if filters and filters.get("card_filter") == "Today":
        order_by_clause = f"CASE WHEN t.priority = 'Urgent' THEN 0 ELSE 1 END, {order_by_clause}"

    return frappe.db.sql(f"""
        SELECT
            t.name,
            t.subject,
            t.status,
            t.priority,
            t.task_owner,
            t.task_owner_name,
            t.task_reporter,
            t.task_reporter_name,
            t.company,
            t.exp_start_date,
            t.exp_end_date,
            t.creation,

            CASE
                WHEN t.status = 'Overdue'
                THEN DATEDIFF(CURDATE(), t.exp_end_date)
                ELSE 0
            END AS overdue_days,

            (
                SELECT c.content
                FROM `tabComment` c
                WHERE c.reference_doctype='Task'
                AND c.reference_name=t.name
                AND c.comment_type='Comment'
                ORDER BY c.creation DESC
                LIMIT 1
            ) AS latest_comment,

            (
                SELECT IFNULL(u.full_name, c.owner)
                FROM `tabComment` c
                LEFT JOIN `tabUser` u ON u.name = c.owner
                WHERE c.reference_doctype='Task'
                AND c.reference_name=t.name
                AND c.comment_type='Comment'
                ORDER BY c.creation DESC
                LIMIT 1
            ) AS latest_comment_by,

            (
                SELECT c.creation
                FROM `tabComment` c
                WHERE c.reference_doctype='Task'
                AND c.reference_name=t.name
                AND c.comment_type='Comment'
                ORDER BY c.creation DESC
                LIMIT 1
            ) AS latest_comment_on

        FROM `tabTask` t
        WHERE {conditions}
        ORDER BY {order_by_clause}
        LIMIT {start}, {page_length}
    """, as_dict=True)


def get_total_count(conditions):
    return frappe.db.sql(f"""
        SELECT COUNT(*) as count
        FROM `tabTask` t
        WHERE {conditions}
    """, as_dict=True)[0].count


@frappe.whitelist()
def get_task_comments(task_name):
    comments = frappe.get_all(
        "Comment",
        filters={
            "reference_doctype": "Task",
            "reference_name": task_name,
            "comment_type": "Comment"
        },
        fields=["name", "content", "owner", "creation"],
        order_by="creation asc"
    )
    
    task = frappe.get_doc("Task", task_name)
    
    return {
        "comments": comments,
        "task_subject": task.subject,
        "task_description": task.description or ""
    }


@frappe.whitelist()
def add_task_comment(task_name, comment):
    if not task_name or not comment:
        frappe.throw(_("Task and Comment are required"))

    doc = frappe.get_doc({
        "doctype": "Comment",
        "comment_type": "Comment",
        "reference_doctype": "Task",
        "reference_name": task_name,
        "content": comment
    })
    doc.insert(ignore_permissions=True)

    return {"status": "success", "message": "Comment added successfully"}


@frappe.whitelist()
def export_tasks(filters=None):
    import openpyxl
    from io import BytesIO
    from frappe.utils.file_manager import save_file

    filters = frappe.parse_json(filters) if filters else {}
    conditions = get_conditions(filters)

    data = frappe.db.sql(f"""
        SELECT
            t.name,
            t.subject,
            t.task_owner_name,
            IFNULL(t.task_reporter_name, t.task_reporter) AS task_reporter,
            t.status,
            t.priority,
            t.creation,
            t.exp_end_date,
            CASE
                WHEN t.status = 'Overdue'
                THEN DATEDIFF(CURDATE(), t.exp_end_date)
                ELSE 0
            END AS overdue_days
        FROM `tabTask` t
        WHERE {conditions}
        ORDER BY t.creation DESC
    """, as_dict=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Task Dashboard"

    ws.append(["Task Dashboard Export"])
    ws.append(["Generated On", getdate(nowdate()).strftime("%d-%m-%Y")])
    ws.append([])

    def get_filter_str(val, default):
        if not val:
            return default
        if isinstance(val, list):
            return ", ".join(str(v) for v in val)
        return val

    ws.append(["Company", get_filter_str(filters.get("company"), "All Companies")])
    ws.append(["Status", get_filter_str(filters.get("status"), "All Statuses")])
    ws.append(["Priority", get_filter_str(filters.get("priority"), "All Priorities")])
    ws.append(["Owner", get_filter_str(filters.get("task_owner"), "All Owners")])
    ws.append(["Search", filters.get("search") or "None"])
    ws.append(["Date From", filters.get("from_date") or "Any"])
    ws.append(["Date To", filters.get("to_date") or "Any"])
    ws.append([])

    headers = [
        "Task ID", "Subject", "Task Owner", "Task Reporter", "Status",
        "Priority", "Created Date", "Due Date", "Overdue Days"
    ]
    ws.append(headers)

    link_font = Font(color="0563C1", underline="single")
    for row in data:
        ws.append([
            row.name,
            row.subject,
            row.task_owner_name,
            row.task_reporter,
            row.status,
            row.priority,
            getdate(row.creation).strftime("%d-%m-%Y") if row.creation else "",
            getdate(row.exp_end_date).strftime("%d-%m-%Y") if row.exp_end_date else "",
            row.overdue_days
        ])
        task_id_cell = ws.cell(row=ws.max_row, column=1)
        task_id_cell.hyperlink = f"{get_url()}/app/task/{row.name}"
        task_id_cell.font = link_font

    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"].font = Font(bold=True)

    header_row_num = 12
    header_fill = PatternFill(start_color="DCE9F7", end_color="DCE9F7", fill_type="solid")
    header_border = Border(bottom=Side(style="thin", color="9DB8D2"))
    for cell in ws[header_row_num]:
        cell.font = Font(bold=True, color="1F3864")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border

    ws.freeze_panes = ws.cell(row=header_row_num + 1, column=1)

    column_widths = [14, 40, 20, 20, 14, 12, 14, 14, 14]
    for idx, width in enumerate(column_widths, start=1):
        ws.column_dimensions[ws.cell(row=header_row_num, column=idx).column_letter].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    file_doc = save_file(
        fname="Task Dashboard Export.xlsx",
        content=output.getvalue(),
        dt=None,
        dn=None,
        is_private=1
    )

    return {
        "file_url": file_doc.file_url
    }

@frappe.whitelist()
def get_leaderboard_data(filters=None, sort_field="completion_rate", sort_order="desc"):
    filters = frappe.parse_json(filters) if filters else {}
    
    base_filters = filters.copy()
    if "card_filter" in base_filters: del base_filters["card_filter"]
    if "status" in base_filters: del base_filters["status"]
        
    conditions = get_conditions(base_filters)

    # Note: Added 'total_count', 'completion_rate', and 'critical_count'
    data = frappe.db.sql(f"""
        SELECT 
            t.task_owner, 
            IFNULL(t.task_owner_name, t.task_owner) as task_owner_name,
            SUM(CASE WHEN t.status = 'Completed' THEN 1 ELSE 0 END) as completed_count,
            SUM(CASE WHEN t.status = 'Overdue' THEN 1 ELSE 0 END) as overdue_count,
            SUM(CASE WHEN t.status NOT IN ('Completed', 'Cancelled', 'Overdue') THEN 1 ELSE 0 END) as open_count,
            COUNT(t.name) as total_count,
            /* Calculation for percentage */
            ROUND((SUM(CASE WHEN t.status = 'Completed' THEN 1 ELSE 0 END) / COUNT(t.name)) * 100, 0) as completion_rate,
            /* Easy to understand Priority: Count of Pending High/Urgent tasks */
            SUM(CASE WHEN t.priority IN ('High', 'Urgent') AND t.status NOT IN ('Completed', 'Cancelled') THEN 1 ELSE 0 END) as critical_count
        FROM `tabTask` t
        WHERE {conditions} AND t.task_owner IS NOT NULL AND t.task_owner != ''
        GROUP BY t.task_owner
        ORDER BY {sort_field} {sort_order}
    """, as_dict=True)

    return { "leaderboard": data }